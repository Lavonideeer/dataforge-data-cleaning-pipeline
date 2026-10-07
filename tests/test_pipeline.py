"""DF-006 end-to-end pipeline, cleaned-output, and CLI tests."""

from __future__ import annotations

import csv
import hashlib
import json
import os
import subprocess
import sys
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

import dataforge.cli as cli_module
import dataforge.pipeline as pipeline_module
from dataforge.audit import AUDIT_COLUMNS, REJECTED_COLUMNS
from dataforge.pipeline import (
    CLEANED_SALES_COLUMNS,
    FINAL_ARTIFACT_NAMES,
    PipelinePathError,
    run_pipeline,
)


ROOT = Path(__file__).resolve().parents[1]
DEMO_RAW = ROOT / "data" / "demo_raw"
DEMO_EXPECTED = ROOT / "data" / "demo_expected"


def _csv(path: Path) -> tuple[tuple[str, ...], tuple[tuple[str, ...], ...]]:
    with path.open(encoding="utf-8", newline="") as stream:
        reader = csv.reader(stream)
        return tuple(next(reader)), tuple(tuple(row) for row in reader)


def _xlsx_text(value: object, column: str) -> str:
    if value is None:
        return ""
    if column == "order_date":
        assert isinstance(value, (date, datetime))
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    if column == "unit_price_eur":
        return format(Decimal(str(value)), ".2f")
    return str(value)


def _xlsx_semantics(path: Path) -> tuple[tuple[str, ...], tuple[tuple[str, ...], ...]]:
    workbook = load_workbook(path, data_only=False)
    try:
        worksheet = workbook["cleaned_sales"]
        values = list(worksheet.iter_rows(values_only=True))
        return (
            tuple(str(value) for value in values[0]),
            tuple(
                tuple(
                    _xlsx_text(value, column)
                    for value, column in zip(row, CLEANED_SALES_COLUMNS, strict=True)
                )
                for row in values[1:]
            ),
        )
    finally:
        workbook.close()


@pytest.fixture(scope="module")
def completed(tmp_path_factory: pytest.TempPathFactory):
    return run_pipeline(DEMO_RAW, tmp_path_factory.mktemp("df006") / "output")


def test_pipeline_result_and_complete_artifact_contract(completed) -> None:
    assert tuple(path.name for path in completed.paths()) == FINAL_ARTIFACT_NAMES
    assert {path.name for path in completed.output_dir.iterdir()} == set(FINAL_ARTIFACT_NAMES)
    assert all(path.is_file() for path in completed.paths())
    assert completed.evidence_outputs.directory == completed.output_dir


def test_corpus_reconciliation_and_evidence_counts(completed) -> None:
    counts = completed.validation.reconciliation
    assert counts.as_dict() == {
        "input": 164,
        "accepted": 146,
        "quarantined": 16,
        "deduplicated": 2,
        "reconciled": True,
    }
    assert len(completed.evidence_outputs.rejected_rows_data) == 23
    assert len(completed.evidence_outputs.audit_events) == 63


def test_cleaned_csv_contract_and_canonical_values(completed) -> None:
    header, rows = _csv(completed.cleaned_sales_csv)
    assert header == CLEANED_SALES_COLUMNS
    assert len(rows) == 146
    assert all(len(row) == len(CLEANED_SALES_COLUMNS) for row in rows)
    assert all(row[3] == date.fromisoformat(row[3]).isoformat() for row in rows)
    assert all(row[4] == str(int(row[4])) for row in rows)
    assert all(row[5] == format(Decimal(row[5]), ".2f") for row in rows)
    assert all(row[9] == str(int(row[9])) for row in rows)
    assert rows[1][:7] == (
        "ORD-0002", "C002", "P012", "2026-07-02", "2", "1299.00", "FR"
    )


def test_cleaned_xlsx_contract_and_csv_equivalence(completed) -> None:
    assert _xlsx_semantics(completed.cleaned_sales_xlsx) == _csv(completed.cleaned_sales_csv)
    workbook = load_workbook(completed.cleaned_sales_xlsx, data_only=False)
    try:
        assert workbook.sheetnames == ["cleaned_sales"]
        worksheet = workbook["cleaned_sales"]
        assert worksheet.sheet_state == "visible"
        assert not any(dimension.hidden for dimension in worksheet.row_dimensions.values())
        assert not any(dimension.hidden for dimension in worksheet.column_dimensions.values())
        assert worksheet.max_row == 147
        assert worksheet.max_column == 10
        assert worksheet.freeze_panes == "A2"
        assert worksheet.auto_filter.ref == "A1:J147"
        assert all(cell.data_type != "f" for row in worksheet.iter_rows() for cell in row)
        assert all(cell.number_format == "yyyy-mm-dd" for cell in worksheet["D"][1:])
        assert all(cell.number_format == "0.00" for cell in worksheet["F"][1:])
        assert all(isinstance(cell.value, int) for cell in worksheet["E"][1:])
        assert all(isinstance(cell.value, int) for cell in worksheet["J"][1:])
    finally:
        workbook.close()


def test_evidence_headers_summary_and_final_report_truth(completed) -> None:
    rejected_header, rejected = _csv(completed.evidence_outputs.rejected_rows)
    audit_header, audit = _csv(completed.evidence_outputs.audit_log)
    summary = json.loads(completed.evidence_outputs.cleaning_summary.read_text(encoding="utf-8"))
    report = completed.evidence_outputs.data_quality_report.read_text(encoding="utf-8")
    assert rejected_header == REJECTED_COLUMNS and len(rejected) == 23
    assert audit_header == AUDIT_COLUMNS and len(audit) == 63
    assert summary["outputs"] == list(FINAL_ARTIFACT_NAMES)
    assert summary["transaction_counts"]["reconciled"] is True
    assert "All six frozen v1 artifacts were generated and verified" in report
    assert "They do not exist yet" not in report
    assert "not produced by this phase" not in report
    for name in FINAL_ARTIFACT_NAMES:
        assert f'href="{name}"' in report


def test_cleaned_output_matches_expected_transaction_oracle(completed) -> None:
    with (DEMO_EXPECTED / "expected_transactions.csv").open(encoding="utf-8", newline="") as stream:
        expected = list(csv.DictReader(stream))
    expected_accepted = {
        (row["source_file"], row["source_sheet"], row["source_row"])
        for row in expected
        if row["expected_terminal_disposition"] == "ACCEPTED"
    }
    expected_excluded = {
        (row["source_file"], row["source_sheet"], row["source_row"])
        for row in expected
        if row["expected_terminal_disposition"] != "ACCEPTED"
    }
    _, cleaned = _csv(completed.cleaned_sales_csv)
    cleaned_refs = {(row[7], row[8], row[9]) for row in cleaned}
    assert cleaned_refs == expected_accepted
    assert not cleaned_refs & expected_excluded
    _, rejected = _csv(completed.evidence_outputs.rejected_rows)
    rejected_sales_refs = {(row[1], row[2], row[3]) for row in rejected if row[0] == "sales"}
    assert rejected_sales_refs == expected_excluded


def test_duplicate_survivor_semantics_reach_cleaned_output(completed) -> None:
    _, rows = _csv(completed.cleaned_sales_csv)
    by_order: dict[str, list[tuple[str, ...]]] = {}
    for row in rows:
        by_order.setdefault(row[0], []).append(row)
    assert all(len(values) == 1 for values in by_order.values())
    assert by_order["ORD-0004"][0][9] == "5"


def test_repeated_runs_are_deterministic(tmp_path: Path) -> None:
    first = run_pipeline(DEMO_RAW, tmp_path / "first")
    second = run_pipeline(DEMO_RAW, tmp_path / "second")
    byte_names = set(FINAL_ARTIFACT_NAMES) - {"cleaned_sales.xlsx"}
    for name in byte_names:
        assert (first.output_dir / name).read_bytes() == (second.output_dir / name).read_bytes()
    assert _xlsx_semantics(first.cleaned_sales_xlsx) == _xlsx_semantics(second.cleaned_sales_xlsx)
    assert first.validation == second.validation


def test_foreign_files_survive_and_known_outputs_are_replaced(tmp_path: Path) -> None:
    output = tmp_path / "output"
    output.mkdir()
    foreign = output / "client_notes.txt"
    foreign.write_text("keep me\n", encoding="utf-8")
    for name in FINAL_ARTIFACT_NAMES:
        (output / name).write_text("stale\n", encoding="utf-8")
    result = run_pipeline(DEMO_RAW, output)
    assert foreign.read_text(encoding="utf-8") == "keep me\n"
    assert all(path.read_bytes() != b"stale\n" for path in result.paths())


@pytest.mark.parametrize("relation", ["same", "nested"])
def test_unsafe_path_relations_fail(relation: str, tmp_path: Path) -> None:
    source = tmp_path / "raw"
    source.mkdir()
    output = source if relation == "same" else source / "output"
    with pytest.raises(PipelinePathError):
        run_pipeline(source, output)


def test_missing_input_and_output_file_fail(tmp_path: Path) -> None:
    with pytest.raises(PipelinePathError, match="does not exist"):
        run_pipeline(tmp_path / "missing", tmp_path / "out")
    output_file = tmp_path / "output.txt"
    output_file.write_text("x", encoding="utf-8")
    with pytest.raises(PipelinePathError, match="not a directory"):
        run_pipeline(DEMO_RAW, output_file)


def test_cli_success_is_concise_and_complete(tmp_path: Path) -> None:
    env = os.environ.copy()
    env["PYTHONPATH"] = os.pathsep.join(
        part for part in (str(ROOT / "src"), env.get("PYTHONPATH", "")) if part
    )
    output = tmp_path / "cli-output"
    completed = subprocess.run(
        [sys.executable, "-m", "dataforge.cli", "--input", str(DEMO_RAW), "--output", str(output)],
        cwd=ROOT,
        env=env,
        text=True,
        capture_output=True,
        check=False,
    )
    assert completed.returncode == 0
    assert completed.stderr == ""
    assert "164 input = 146 accepted + 16 quarantined + 2 deduplicated" in completed.stdout
    assert all(name in completed.stdout for name in FINAL_ARTIFACT_NAMES)
    assert len(completed.stdout.splitlines()) == 3
    assert {path.name for path in output.iterdir()} == set(FINAL_ARTIFACT_NAMES)


def test_cli_expected_failures_have_no_traceback(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    assert cli_module.main(["--input", str(tmp_path / "missing"), "--output", str(tmp_path / "out")]) != 0
    captured = capsys.readouterr()
    assert "DataForge failed:" in captured.err
    assert "Traceback" not in captured.err
    assert captured.out == ""


def test_cli_malformed_schema_is_nonzero(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    source = tmp_path / "malformed"
    source.mkdir()
    (source / "sales.csv").write_text("order_id,quantity\nORD-1,1\n", encoding="utf-8")
    assert cli_module.main(["--input", str(source), "--output", str(tmp_path / "out")]) != 0
    captured = capsys.readouterr()
    assert "MISSING_REQUIRED_COLUMN" in captured.err
    assert "Traceback" not in captured.err
    assert not (tmp_path / "out").exists()


def test_cli_unreadable_workbook_is_nonzero(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    source = tmp_path / "broken-workbook"
    source.mkdir()
    (source / "customers.xlsx").write_bytes(b"not an xlsx file")
    assert cli_module.main(["--input", str(source), "--output", str(tmp_path / "out")]) != 0
    captured = capsys.readouterr()
    assert "DataForge failed:" in captured.err
    assert "Traceback" not in captured.err
    assert not (tmp_path / "out").exists()


def test_pipeline_production_is_fixture_independent_and_count_independent() -> None:
    for module in (pipeline_module, cli_module):
        source = Path(module.__file__).read_text(encoding="utf-8")
        for forbidden in ("demo_expected", "expected_transactions", "expected_outcomes", "ground_truth_"):
            assert forbidden not in source
    pipeline_source = Path(pipeline_module.__file__).read_text(encoding="utf-8")
    for fixture_count in ("146", "164", "23", "63"):
        assert fixture_count not in pipeline_source


def test_artifact_sizes_are_nonzero_and_csv_has_lf_endings(completed) -> None:
    assert all(path.stat().st_size > 0 for path in completed.paths())
    for name in ("cleaned_sales.csv", "rejected_rows.csv", "audit_log.csv"):
        payload = (completed.output_dir / name).read_bytes()
        assert b"\r\n" not in payload
        assert payload.endswith(b"\n")


def test_generated_file_digests_are_stable_except_xlsx(tmp_path: Path) -> None:
    first = run_pipeline(DEMO_RAW, tmp_path / "a")
    second = run_pipeline(DEMO_RAW, tmp_path / "b")
    names = set(FINAL_ARTIFACT_NAMES) - {"cleaned_sales.xlsx"}
    digest = lambda path: hashlib.sha256(path.read_bytes()).hexdigest()
    assert {name: digest(first.output_dir / name) for name in names} == {
        name: digest(second.output_dir / name) for name in names
    }
