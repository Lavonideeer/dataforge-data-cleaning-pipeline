"""DF-001 semantic tests for the deterministic synthetic corpus."""

from __future__ import annotations

import csv
import hashlib
import json
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from scripts.generate_demo_data import DEMO_DATE, EXPECTED_FILES, RAW_FILES, SEED, generate_demo_data


ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data" / "demo_raw"
EXPECTED = ROOT / "data" / "demo_expected"


def _json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def _csv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8", newline="") as stream:
        return list(csv.DictReader(stream))


def _xlsx(path: Path) -> list[dict]:
    workbook = load_workbook(path, data_only=False, read_only=True)
    result = [{"sheet": sheet.title, "rows": [[cell.value for cell in row] for row in sheet.iter_rows()]} for sheet in workbook.worksheets]
    workbook.close()
    return result


def _source_row(source: dict) -> list:
    path = RAW / source["source_file"]
    if path.suffix == ".csv":
        with path.open(encoding="utf-8", newline="") as stream:
            return list(csv.reader(stream))[source["source_row"] - 1]
    workbook = load_workbook(path, data_only=False, read_only=True)
    values = [cell.value for cell in workbook[source["source_sheet"]][source["source_row"]]]
    workbook.close()
    return values


def test_generation_is_byte_or_semantically_deterministic(tmp_path: Path) -> None:
    first_root, second_root = tmp_path / "first", tmp_path / "second"
    assert generate_demo_data(first_root) == generate_demo_data(second_root)
    for filename in RAW_FILES + EXPECTED_FILES:
        subdir = "demo_raw" if filename in RAW_FILES else "demo_expected"
        first, second = first_root / "data" / subdir / filename, second_root / "data" / subdir / filename
        assert _xlsx(first) == _xlsx(second) if first.suffix == ".xlsx" else first.read_bytes() == second.read_bytes()


def test_manifest_covers_every_frozen_defect_class() -> None:
    manifest = _json(EXPECTED / "expected_outcomes.json")
    required = {"schema.documented_aliases", "schema.unexpected_columns", "locale.currency_formats", "locale.currency_invalid", "dates.iso", "dates.ddmmyyyy", "dates.textual", "dates.impossible", "dates.future", "country.aliases", "country.unknown", "categories.aliases", "categories.unknown", "missingness.critical_identifier", "missingness.optional_field", "duplicates.exact_row", "duplicates.normalized_business_key", "duplicates.conflicting_business_key", "duplicates.fuzzy_customer", "business.quantity_nonpositive", "business.negative_price", "text.whitespace", "text.case", "email.hygiene", "email.malformed", "references.valid", "references.unknown_customer", "references.unknown_product", "precedence.hard_before_duplicate", "failures.multiple"}
    assert required <= set(manifest["coverage"])
    known = {case["case_id"] for case in manifest["cases"]}
    assert all(case_id in known for key in required for case_id in manifest["coverage"][key])


def test_cases_a_through_j_exist_and_reconcile() -> None:
    manifest = _json(EXPECTED / "expected_outcomes.json")
    assert {case["interaction_case"] for case in manifest["cases"] if case["interaction_case"]} == set("ABCDEFGHIJ")
    g2 = manifest["g2_expected"]
    assert g2["input_transaction_rows"] == g2["accepted_rows"] + g2["quarantined_rows"] + g2["deduplicated_rows"]
    assert g2 == {"accepted_rows": 146, "deduplicated_rows": 2, "input_transaction_rows": 164, "quarantined_rows": 16, "reconciled": True}


def test_frozen_clock_and_date_fixtures() -> None:
    assert SEED == 1007 and DEMO_DATE == date(2026, 10, 7)
    manifest = _json(EXPECTED / "expected_outcomes.json")
    transactions = {row["fixture_id"]: row for row in _csv(EXPECTED / "expected_transactions.csv")}
    for case in [case for case in manifest["cases"] if case["interaction_case"] in {"I", "J"}]:
        for source in case["source_records"]:
            if source.get("record_type") != "sales":
                continue
            raw = _source_row(source)
            path = RAW / source["source_file"]
            if path.suffix == ".xlsx":
                workbook = load_workbook(path, read_only=True, data_only=True)
                headers = [cell.value for cell in workbook[source["source_sheet"]][1]]
                workbook.close()
            else:
                with path.open(encoding="utf-8", newline="") as stream:
                    headers = next(csv.reader(stream))
            date_index = next(index for index, value in enumerate(headers) if value in {"order_date", "Order Date", "OrderDate", "order date", "Date"})
            value = str(raw[date_index])
            parsed = datetime.strptime(value, "%d/%m/%Y" if "/" in value else "%Y-%m-%d").date()
            assert parsed > DEMO_DATE
            assert transactions[source["fixture_id"]]["expected_terminal_disposition"] == "QUARANTINED"
    assert all(datetime.strptime(row["order_date"], "%Y-%m-%d").date() <= DEMO_DATE for row in _csv(EXPECTED / "ground_truth_sales.csv"))


def test_provenance_identities_resolve_to_unique_raw_rows() -> None:
    rows = _csv(EXPECTED / "expected_transactions.csv")
    identities = {(row["source_file"], row["source_sheet"], int(row["source_row"])) for row in rows}
    assert len(rows) == len(identities) == 164
    assert all(_source_row({"source_file": row["source_file"], "source_sheet": row["source_sheet"], "source_row": int(row["source_row"])}) for row in rows)


def test_duplicate_fixtures_match_manifest_claims() -> None:
    cases = {case["case_id"]: case for case in _json(EXPECTED / "expected_outcomes.json")["cases"]}
    exact = cases["D_EXACT_DUPLICATE"]["source_records"]
    assert len(exact) == 2 and _source_row(exact[0]) == _source_row(exact[1])
    conflict = cases["E_CONFLICTING_ORDER_ID"]["source_records"]
    values = [_source_row(source) for source in conflict]
    assert len(values) == 2 and values[0][0] == values[1][0] and values[0][4] != values[1][4]
    assert set(cases["E_CONFLICTING_ORDER_ID"]["expected_terminal_dispositions"]) == {"QUARANTINED"}
    assert sorted(cases["F_HARD_INVALID_BEFORE_DUPLICATE"]["expected_terminal_dispositions"]) == ["ACCEPTED", "QUARANTINED"]


def test_outputs_are_readable_and_synthetic() -> None:
    assert {path.suffix for path in RAW.glob("sales_*")} == {".csv", ".xlsx"}
    for filename in RAW_FILES:
        path = RAW / filename
        assert path.is_file() and path.stat().st_size > 0
        if path.suffix == ".xlsx":
            workbook = load_workbook(path, read_only=True, data_only=True)
            assert workbook.sheetnames
            workbook.close()
        else:
            assert _csv(path)
    workbook = load_workbook(RAW / "customers.xlsx", read_only=True, data_only=True)
    rows = list(workbook["customers"].iter_rows(values_only=True))
    workbook.close()
    email_index = rows[0].index("Email Address")
    emails = [str(row[email_index]).strip().lower() for row in rows[1:] if row[email_index]]
    assert all("@" not in value or value.endswith("@example.test") for value in emails)
    assert _json(EXPECTED / "expected_outcomes.json")["dataset"]["contains_real_customer_data"] is False


def test_committed_manifest_integrity_markers_match_files() -> None:
    for entry in _json(EXPECTED / "expected_outcomes.json")["files"]:
        path = ROOT / entry["path"]
        if entry["comparison"] == "byte":
            assert hashlib.sha256(path.read_bytes()).hexdigest() == entry["sha256"]
        else:
            payload = json.dumps(_xlsx(path), ensure_ascii=False, sort_keys=True, separators=(",", ":"), default=str).encode()
            assert hashlib.sha256(payload).hexdigest() == entry["semantic_sha256"]
