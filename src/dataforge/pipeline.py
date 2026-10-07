"""Deterministic DF-006 orchestration and cleaned-sales delivery.

The pipeline composes the existing DF-002 through DF-005 interfaces. It does
not re-decide business facts. Its responsibilities are bounded to path safety,
canonical serialization, output verification, and fail-safe publication.
"""

from __future__ import annotations

import csv
import os
import tempfile
from dataclasses import dataclass, replace
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Mapping, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from .audit import V1_OUTPUT_NAMES
from .ingest import IngestionResult, ingest_directory
from .normalize import NormalizationResult, normalize
from .report import EvidenceOutputs, write_evidence_outputs
from .validate import ACCEPTED, ValidationResult, validate


CLEANED_SALES_COLUMNS: tuple[str, ...] = (
    "order_id",
    "customer_id",
    "product_id",
    "order_date",
    "quantity",
    "unit_price_eur",
    "country",
    "_source_file",
    "_source_sheet",
    "_source_row",
)

FINAL_ARTIFACT_NAMES: tuple[str, ...] = V1_OUTPUT_NAMES
_FIXED_WORKBOOK_TIME = datetime(2026, 10, 7, 0, 0, 0)


class PipelineError(ValueError):
    """Expected, actionable failure of the bounded client workflow."""


class PipelinePathError(PipelineError):
    """Unsafe or unusable input/output path relation."""


class PipelineInvariantError(PipelineError):
    """Upstream results or staged artifacts violate the frozen contract."""


@dataclass(frozen=True)
class PipelineResult:
    """Structured result of one successfully published DF-006 run."""

    input_dir: Path
    output_dir: Path
    ingestion: IngestionResult
    normalization: NormalizationResult
    validation: ValidationResult
    evidence_outputs: EvidenceOutputs
    cleaned_sales_csv: Path
    cleaned_sales_xlsx: Path

    def paths(self) -> tuple[Path, ...]:
        """Return all six artifacts in the frozen appendix order."""

        by_name = {
            self.cleaned_sales_xlsx.name: self.cleaned_sales_xlsx,
            self.cleaned_sales_csv.name: self.cleaned_sales_csv,
        }
        by_name.update({path.name: path for path in self.evidence_outputs.paths()})
        return tuple(by_name[name] for name in FINAL_ARTIFACT_NAMES)


def _resolved_paths(input_dir: str | Path, output_dir: str | Path) -> tuple[Path, Path]:
    source = Path(input_dir)
    destination = Path(output_dir)
    if not source.exists():
        raise PipelinePathError(f"Input directory does not exist: {source}")
    if not source.is_dir():
        raise PipelinePathError(f"Input path is not a directory: {source}")
    if destination.exists() and not destination.is_dir():
        raise PipelinePathError(f"Output path is not a directory: {destination}")

    source_resolved = source.resolve()
    destination_resolved = destination.resolve(strict=False)
    if source_resolved == destination_resolved:
        raise PipelinePathError("Input and output directories must be different.")
    if destination_resolved.is_relative_to(source_resolved):
        raise PipelinePathError("Output directory must not be inside the raw input directory.")
    return source_resolved, destination_resolved


def _assert_input_roles(ingestion: IngestionResult) -> None:
    roles = {item.record_type for item in ingestion.file_metadata}
    missing = [name for name in ("sales", "customers", "products") if name not in roles]
    if missing:
        raise PipelineInvariantError(
            "Input directory is missing required bounded input role(s): " + ", ".join(missing)
        )


def _serialize_cleaned_value(column: str, value: object) -> str:
    """Serialize one already-normalized value without applying a new business rule."""

    if value is None:
        return ""
    if column in {"quantity", "_source_row"}:
        if isinstance(value, bool):
            raise PipelineInvariantError(f"Canonical {column} is boolean, not an integer.")
        try:
            integer = int(value)
        except (TypeError, ValueError) as exc:
            raise PipelineInvariantError(f"Canonical {column} is not an integer: {value!r}") from exc
        if str(integer) != str(value):
            raise PipelineInvariantError(f"Canonical {column} is not exact integer text: {value!r}")
        return str(integer)
    if column == "unit_price_eur":
        try:
            amount = value if isinstance(value, Decimal) else Decimal(str(value))
        except (InvalidOperation, ValueError) as exc:
            raise PipelineInvariantError(f"Canonical money is invalid: {value!r}") from exc
        if not amount.is_finite() or amount.as_tuple().exponent < -2:
            raise PipelineInvariantError(f"Canonical money is not fixed precision: {value!r}")
        return format(amount, ".2f")
    if column == "order_date":
        if isinstance(value, datetime):
            if value.time() != time.min or value.tzinfo is not None:
                raise PipelineInvariantError(f"Canonical date contains a timestamp: {value!r}")
            return value.date().isoformat()
        if isinstance(value, date):
            return value.isoformat()
        try:
            return date.fromisoformat(str(value)).isoformat()
        except ValueError as exc:
            raise PipelineInvariantError(f"Canonical date is invalid: {value!r}") from exc
    return str(value)


def _cleaned_rows(validation: ValidationResult) -> tuple[tuple[str, ...], ...]:
    if not validation.reconciliation.reconciled:
        raise PipelineInvariantError("G2 reconciliation is not satisfied.")
    if len(validation.accepted) != validation.reconciliation.accepted_rows:
        raise PipelineInvariantError("Accepted-row bucket disagrees with reconciliation.")

    refs = [outcome.ref for outcome in validation.accepted]
    if len(refs) != len(set(refs)):
        raise PipelineInvariantError("An accepted transaction identity appears more than once.")
    rows: list[tuple[str, ...]] = []
    for outcome in validation.accepted:
        if outcome.disposition != ACCEPTED:
            raise PipelineInvariantError("Cleaned output received a non-accepted transaction.")
        values = dict(outcome.values)
        values.update(
            {
                "_source_file": outcome.ref.source_file,
                "_source_sheet": outcome.ref.source_sheet,
                "_source_row": outcome.ref.source_row,
            }
        )
        rows.append(
            tuple(
                _serialize_cleaned_value(column, values.get(column))
                for column in CLEANED_SALES_COLUMNS
            )
        )
    return tuple(rows)


def _write_cleaned_csv(path: Path, rows: Sequence[Sequence[str]]) -> Path:
    with path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.writer(stream, lineterminator="\n")
        writer.writerow(CLEANED_SALES_COLUMNS)
        writer.writerows(rows)
    return path


def _xlsx_value(column: str, text: str) -> object:
    if text == "":
        return None
    if column == "order_date":
        return date.fromisoformat(text)
    if column in {"quantity", "_source_row"}:
        return int(text)
    if column == "unit_price_eur":
        return Decimal(text)
    return text


def _write_cleaned_xlsx(path: Path, rows: Sequence[Sequence[str]]) -> Path:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "cleaned_sales"
    worksheet.append(CLEANED_SALES_COLUMNS)
    for row in rows:
        worksheet.append(
            tuple(
                _xlsx_value(column, value)
                for column, value in zip(CLEANED_SALES_COLUMNS, row, strict=True)
            )
        )

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:J{len(rows) + 1}"
    for cell in worksheet[1]:
        cell.font = Font(bold=True)
    for column, width in {
        "A": 14,
        "B": 14,
        "C": 13,
        "D": 13,
        "E": 10,
        "F": 16,
        "G": 12,
        "H": 24,
        "I": 16,
        "J": 13,
    }.items():
        worksheet.column_dimensions[column].width = width
    for cell in worksheet["D"][1:]:
        cell.number_format = "yyyy-mm-dd"
    for cell in worksheet["F"][1:]:
        cell.number_format = "0.00"
    for cell in worksheet["E"][1:] + worksheet["J"][1:]:
        cell.number_format = "0"

    workbook.properties.created = _FIXED_WORKBOOK_TIME
    workbook.properties.modified = _FIXED_WORKBOOK_TIME
    workbook.save(path)
    workbook.close()
    return path


def _read_csv_rows(path: Path) -> tuple[tuple[str, ...], tuple[tuple[str, ...], ...]]:
    with path.open("r", encoding="utf-8", newline="") as stream:
        reader = csv.reader(stream)
        header = tuple(next(reader, ()))
        rows = tuple(tuple(row) for row in reader)
    return header, rows


def _read_xlsx_rows(path: Path) -> tuple[tuple[str, ...], tuple[tuple[str, ...], ...]]:
    workbook = load_workbook(path, read_only=False, data_only=False)
    try:
        if workbook.sheetnames != ["cleaned_sales"]:
            raise PipelineInvariantError("Cleaned workbook must contain only the cleaned_sales sheet.")
        worksheet = workbook["cleaned_sales"]
        if worksheet.sheet_state != "visible":
            raise PipelineInvariantError("Cleaned workbook data sheet must be visible.")
        if any(dimension.hidden for dimension in worksheet.row_dimensions.values()):
            raise PipelineInvariantError("Cleaned workbook must not contain hidden rows.")
        if any(dimension.hidden for dimension in worksheet.column_dimensions.values()):
            raise PipelineInvariantError("Cleaned workbook must not contain hidden columns.")
        if worksheet.freeze_panes != "A2":
            raise PipelineInvariantError("Cleaned workbook must freeze the header row.")
        expected_filter = f"A1:J{worksheet.max_row}"
        if worksheet.auto_filter.ref != expected_filter:
            raise PipelineInvariantError("Cleaned workbook autofilter does not cover the data table.")
        if any(cell.data_type == "f" for row in worksheet.iter_rows() for cell in row):
            raise PipelineInvariantError("Cleaned workbook must not contain formulas.")

        values = list(worksheet.iter_rows(values_only=True))
        header = tuple("" if value is None else str(value) for value in values[0]) if values else ()
        rows = tuple(
            tuple(
                _serialize_cleaned_value(column, value)
                for column, value in zip(CLEANED_SALES_COLUMNS, row, strict=True)
            )
            for row in values[1:]
        )
        return header, rows
    finally:
        workbook.close()


def _verify_staged_delivery(
    staging: Path,
    rows: tuple[tuple[str, ...], ...],
    validation: ValidationResult,
    evidence: EvidenceOutputs,
) -> None:
    produced = {path.name for path in staging.iterdir() if path.is_file()}
    if produced != set(FINAL_ARTIFACT_NAMES):
        raise PipelineInvariantError(f"Staged artifact contract mismatch: {sorted(produced)!r}")

    csv_header, csv_rows = _read_csv_rows(staging / "cleaned_sales.csv")
    xlsx_header, xlsx_rows = _read_xlsx_rows(staging / "cleaned_sales.xlsx")
    if csv_header != CLEANED_SALES_COLUMNS or xlsx_header != CLEANED_SALES_COLUMNS:
        raise PipelineInvariantError("Cleaned output header/order violates the frozen schema.")
    if csv_rows != rows or xlsx_rows != rows or csv_rows != xlsx_rows:
        raise PipelineInvariantError("Cleaned CSV/XLSX values or ordering disagree.")
    if len(rows) != validation.reconciliation.accepted_rows:
        raise PipelineInvariantError("Cleaned row count disagrees with accepted reconciliation count.")

    counts = evidence.summary.get("transaction_counts")
    if not isinstance(counts, Mapping) or counts != validation.reconciliation.as_dict():
        raise PipelineInvariantError("Evidence summary disagrees with validation reconciliation.")
    if evidence.summary.get("outputs") != list(FINAL_ARTIFACT_NAMES):
        raise PipelineInvariantError("Evidence summary output list violates the six-artifact contract.")
    report = evidence.data_quality_report.read_text(encoding="utf-8")
    if "They do not exist yet" in report or "not produced by this phase" in report:
        raise PipelineInvariantError("Final HTML incorrectly describes cleaned outputs as pending.")


def _publish(staging: Path, output_dir: Path) -> None:
    """Publish only the six known files, restoring prior versions on failure."""

    created_directory = not output_dir.exists()
    output_dir.mkdir(parents=True, exist_ok=True)
    backup = staging / ".previous"
    backup.mkdir()
    moved: list[str] = []
    published: list[str] = []
    try:
        for name in FINAL_ARTIFACT_NAMES:
            target = output_dir / name
            if target.exists() or target.is_symlink():
                if target.is_dir():
                    raise PipelinePathError(f"Known output target is a directory: {target}")
                os.replace(target, backup / name)
                moved.append(name)
        for name in FINAL_ARTIFACT_NAMES:
            os.replace(staging / name, output_dir / name)
            published.append(name)
    except Exception:
        for name in published:
            target = output_dir / name
            if target.exists() or target.is_symlink():
                target.unlink()
        for name in moved:
            os.replace(backup / name, output_dir / name)
        if created_directory:
            try:
                output_dir.rmdir()
            except OSError:
                pass
        raise


def run_pipeline(input_dir: str | Path, output_dir: str | Path) -> PipelineResult:
    """Run DF-002 through DF-006 and publish the verified six-file delivery."""

    source, destination = _resolved_paths(input_dir, output_dir)
    destination.parent.mkdir(parents=True, exist_ok=True)

    ingestion = ingest_directory(source)
    _assert_input_roles(ingestion)
    normalization = normalize(ingestion)
    validation = validate(normalization)
    rows = _cleaned_rows(validation)

    with tempfile.TemporaryDirectory(prefix=".dataforge-staging-", dir=destination.parent) as temp:
        staging = Path(temp)
        _write_cleaned_csv(staging / "cleaned_sales.csv", rows)
        _write_cleaned_xlsx(staging / "cleaned_sales.xlsx", rows)
        evidence = write_evidence_outputs(
            normalization,
            validation,
            staging,
            ingestion=ingestion,
            input_dir=source,
            full_pipeline=True,
        )
        _verify_staged_delivery(staging, rows, validation, evidence)
        _publish(staging, destination)

    published_evidence = replace(
        evidence,
        directory=destination,
        rejected_rows=destination / "rejected_rows.csv",
        audit_log=destination / "audit_log.csv",
        cleaning_summary=destination / "cleaning_summary.json",
        data_quality_report=destination / "data_quality_report.html",
    )
    return PipelineResult(
        input_dir=source,
        output_dir=destination,
        ingestion=ingestion,
        normalization=normalization,
        validation=validation,
        evidence_outputs=published_evidence,
        cleaned_sales_csv=destination / "cleaned_sales.csv",
        cleaned_sales_xlsx=destination / "cleaned_sales.xlsx",
    )


__all__ = [
    "CLEANED_SALES_COLUMNS",
    "FINAL_ARTIFACT_NAMES",
    "PipelineError",
    "PipelineInvariantError",
    "PipelinePathError",
    "PipelineResult",
    "run_pipeline",
]
