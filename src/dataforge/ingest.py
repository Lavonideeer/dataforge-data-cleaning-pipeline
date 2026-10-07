"""Deterministic CSV/XLSX ingestion for the bounded DataForge demo."""

from __future__ import annotations

import csv
import unicodedata
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, Mapping

from openpyxl import load_workbook

from .schema import PROVENANCE_FIELDS, RawValue, RecordType, SCHEMAS, SchemaDiagnostic, map_headers


SUPPORTED_SUFFIXES = frozenset({".csv", ".xlsx"})


@dataclass(frozen=True)
class DiscoveredInput:
    path: Path
    source_file: str
    record_type: RecordType
    selected_sheet: str


@dataclass(frozen=True)
class IngestedTable:
    record_type: RecordType
    source_file: str
    source_sheet: str
    columns: tuple[str, ...]
    rows: tuple[Mapping[str, RawValue], ...]
    raw_headers: tuple[str, ...]
    aliases_applied: int


@dataclass(frozen=True)
class FileMetadata:
    source_file: str
    record_type: RecordType
    file_format: str
    sheets_inspected: tuple[str, ...]
    selected_sheet: str
    row_count: int
    aliases_applied: int
    unexpected_columns: int


@dataclass(frozen=True)
class IngestionResult:
    discovered_files: tuple[str, ...]
    tables: tuple[IngestedTable, ...]
    diagnostics: tuple[SchemaDiagnostic, ...]
    file_metadata: tuple[FileMetadata, ...]

    def rows_for(self, record_type: RecordType) -> tuple[Mapping[str, RawValue], ...]:
        return tuple(
            row
            for table in self.tables
            if table.record_type == record_type
            for row in table.rows
        )

    @property
    def aliases_applied(self) -> int:
        return sum(item.aliases_applied for item in self.file_metadata)

    @property
    def sheets_inspected(self) -> int:
        return sum(len(item.sheets_inspected) for item in self.file_metadata)

    @property
    def has_errors(self) -> bool:
        return any(item.severity == "ERROR" for item in self.diagnostics)

    def summary(self) -> dict[str, int]:
        return {
            "files_discovered": len(self.discovered_files),
            "files_classified": len(self.file_metadata),
            "sheets_inspected": self.sheets_inspected,
            "sales_rows": len(self.rows_for("sales")),
            "customer_rows": len(self.rows_for("customers")),
            "product_rows": len(self.rows_for("products")),
            "aliases_applied": self.aliases_applied,
            "unexpected_columns": sum(
                item.code == "UNEXPECTED_COLUMN" for item in self.diagnostics
            ),
            "unexpected_sheets": sum(
                item.code == "UNEXPECTED_SHEET" for item in self.diagnostics
            ),
            "unexpected_inputs": sum(
                item.code == "UNEXPECTED_INPUT" for item in self.diagnostics
            ),
            "structural_errors": sum(item.severity == "ERROR" for item in self.diagnostics),
        }


class SchemaIngestionError(ValueError):
    """A file-level schema failure; no row tables are exposed downstream."""

    def __init__(
        self,
        diagnostics: Iterable[SchemaDiagnostic],
        *,
        discovered_files: Iterable[str] = (),
        file_metadata: Iterable[FileMetadata] = (),
    ) -> None:
        self.diagnostics = tuple(diagnostics)
        self.discovered_files = tuple(discovered_files)
        self.file_metadata = tuple(file_metadata)
        codes = ", ".join(item.code for item in self.diagnostics if item.severity == "ERROR")
        super().__init__(f"DataForge schema ingestion failed: {codes or 'structural error'}")


def _source_name(name: str) -> str:
    return unicodedata.normalize("NFC", name).replace("\\", "/")


def _source_sort_key(source_file: str) -> bytes:
    return source_file.encode("utf-8")


def classify_input(filename: str | Path) -> tuple[RecordType, str]:
    """Classify one supported basename under the exhaustive input contract."""

    basename = Path(filename).name
    suffix = Path(basename).suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise ValueError(f"Unsupported input type: {basename}")
    if basename == "customers.xlsx":
        return "customers", "customers"
    if basename == "products.csv":
        return "products", ""
    return "sales", "sales" if suffix == ".xlsx" else ""


def discover_input_files(input_dir: str | Path) -> tuple[DiscoveredInput, ...]:
    """Discover immediate supported regular files in deterministic source order."""

    root = Path(input_dir)
    if not root.is_dir():
        raise NotADirectoryError(f"Input root is not a directory: {root}")
    root_resolved = root.resolve()
    discovered: list[DiscoveredInput] = []
    normalized_names: set[str] = set()
    for path in root.iterdir():
        if path.is_symlink() or not path.is_file() or path.suffix.lower() not in SUPPORTED_SUFFIXES:
            continue
        if not path.resolve().is_relative_to(root_resolved):
            continue
        source_file = _source_name(path.relative_to(root).as_posix())
        if source_file in normalized_names:
            raise SchemaIngestionError(
                (
                    SchemaDiagnostic(
                        code="UNEXPECTED_INPUT",
                        severity="ERROR",
                        record_type=None,
                        source_file=source_file,
                        message="Input paths collide after Unicode-NFC normalization.",
                    ),
                )
            )
        normalized_names.add(source_file)
        record_type, selected_sheet = classify_input(path.name)
        discovered.append(DiscoveredInput(path, source_file, record_type, selected_sheet))
    return tuple(sorted(discovered, key=lambda item: _source_sort_key(item.source_file)))


def _unexpected_input_diagnostics(input_dir: Path) -> list[SchemaDiagnostic]:
    diagnostics: list[SchemaDiagnostic] = []
    for path in input_dir.iterdir():
        if not path.is_file() and not path.is_symlink():
            continue
        if not path.is_symlink() and path.suffix.lower() in SUPPORTED_SUFFIXES:
            continue
        source_file = _source_name(path.relative_to(input_dir).as_posix())
        diagnostics.append(
            SchemaDiagnostic(
                code="UNEXPECTED_INPUT",
                severity="WARNING",
                record_type=None,
                source_file=source_file,
                message=f"Unsupported input entry {source_file!r} ignored.",
            )
        )
    return sorted(diagnostics, key=lambda item: _source_sort_key(item.source_file))


def _row_from_values(
    values: Iterable[RawValue],
    canonical_by_index: tuple[str | None, ...],
    fields: tuple[str, ...],
    *,
    source_file: str,
    source_sheet: str,
    source_row: int,
) -> dict[str, RawValue]:
    cells = tuple(values)
    row: dict[str, RawValue] = {}
    for field in fields:
        index = canonical_by_index.index(field)
        row[field] = cells[index] if index < len(cells) else ""
    row["_source_file"] = source_file
    row["_source_sheet"] = source_sheet
    row["_source_row"] = source_row
    return row


def _read_csv(item: DiscoveredInput) -> tuple[IngestedTable, FileMetadata, tuple[SchemaDiagnostic, ...]]:
    schema = SCHEMAS[item.record_type]
    with item.path.open("r", encoding="utf-8-sig", newline="") as stream:
        reader = csv.reader(stream)
        headers = next(reader, [])
        mapping = map_headers(headers, schema, source_file=item.source_file)
        rows = tuple(
            _row_from_values(
                values,
                mapping.canonical_by_index,
                schema.fields,
                source_file=item.source_file,
                source_sheet="",
                source_row=source_row,
            )
            for source_row, values in enumerate(reader, 2)
        ) if not mapping.has_errors else ()
    diagnostics = mapping.diagnostics
    metadata = FileMetadata(
        source_file=item.source_file,
        record_type=item.record_type,
        file_format="csv",
        sheets_inspected=(),
        selected_sheet="",
        row_count=len(rows),
        aliases_applied=mapping.aliases_applied,
        unexpected_columns=sum(item.code == "UNEXPECTED_COLUMN" for item in diagnostics),
    )
    table = IngestedTable(
        record_type=item.record_type,
        source_file=item.source_file,
        source_sheet="",
        columns=schema.fields + PROVENANCE_FIELDS,
        rows=rows,
        raw_headers=mapping.raw_headers,
        aliases_applied=mapping.aliases_applied,
    )
    return table, metadata, diagnostics


def _read_xlsx(item: DiscoveredInput) -> tuple[IngestedTable | None, FileMetadata, tuple[SchemaDiagnostic, ...]]:
    schema = SCHEMAS[item.record_type]
    workbook = load_workbook(item.path, read_only=True, data_only=False)
    try:
        sheet_names = tuple(workbook.sheetnames)
        diagnostics = [
            SchemaDiagnostic(
                code="UNEXPECTED_SHEET",
                severity="WARNING",
                record_type=item.record_type,
                source_file=item.source_file,
                source_sheet=sheet_name,
                message=f"Unselected worksheet {sheet_name!r} ignored.",
            )
            for sheet_name in sheet_names
            if sheet_name != item.selected_sheet
        ]
        if item.selected_sheet not in sheet_names:
            diagnostics.append(
                SchemaDiagnostic(
                    code="UNEXPECTED_SHEET",
                    severity="ERROR",
                    record_type=item.record_type,
                    source_file=item.source_file,
                    source_sheet=item.selected_sheet,
                    message=f"Required worksheet {item.selected_sheet!r} is missing.",
                )
            )
            metadata = FileMetadata(
                item.source_file,
                item.record_type,
                "xlsx",
                sheet_names,
                item.selected_sheet,
                0,
                0,
                0,
            )
            return None, metadata, tuple(diagnostics)

        worksheet = workbook[item.selected_sheet]
        values = worksheet.iter_rows(values_only=True)
        headers = next(values, ())
        mapping = map_headers(
            headers,
            schema,
            source_file=item.source_file,
            source_sheet=item.selected_sheet,
        )
        diagnostics.extend(mapping.diagnostics)
        rows = tuple(
            _row_from_values(
                row_values,
                mapping.canonical_by_index,
                schema.fields,
                source_file=item.source_file,
                source_sheet=item.selected_sheet,
                source_row=source_row,
            )
            for source_row, row_values in enumerate(values, 2)
        ) if not mapping.has_errors else ()
        metadata = FileMetadata(
            source_file=item.source_file,
            record_type=item.record_type,
            file_format="xlsx",
            sheets_inspected=sheet_names,
            selected_sheet=item.selected_sheet,
            row_count=len(rows),
            aliases_applied=mapping.aliases_applied,
            unexpected_columns=sum(value.code == "UNEXPECTED_COLUMN" for value in diagnostics),
        )
        table = IngestedTable(
            record_type=item.record_type,
            source_file=item.source_file,
            source_sheet=item.selected_sheet,
            columns=schema.fields + PROVENANCE_FIELDS,
            rows=rows,
            raw_headers=mapping.raw_headers,
            aliases_applied=mapping.aliases_applied,
        )
        return table, metadata, tuple(diagnostics)
    finally:
        workbook.close()


def ingest_directory(input_dir: str | Path) -> IngestionResult:
    """Ingest one bounded input root, failing atomically on schema errors."""

    root = Path(input_dir)
    discovered = discover_input_files(root)
    diagnostics = _unexpected_input_diagnostics(root)
    tables: list[IngestedTable] = []
    metadata: list[FileMetadata] = []
    for item in discovered:
        table, file_metadata, file_diagnostics = (
            _read_csv(item) if item.path.suffix.lower() == ".csv" else _read_xlsx(item)
        )
        metadata.append(file_metadata)
        diagnostics.extend(file_diagnostics)
        if table is not None:
            tables.append(table)

    if any(item.severity == "ERROR" for item in diagnostics):
        raise SchemaIngestionError(
            diagnostics,
            discovered_files=(item.source_file for item in discovered),
            file_metadata=metadata,
        )
    return IngestionResult(
        discovered_files=tuple(item.source_file for item in discovered),
        tables=tuple(tables),
        diagnostics=tuple(diagnostics),
        file_metadata=tuple(metadata),
    )
