"""Generate the deterministic DF-001 synthetic demo corpus.

Canonical ground truth is created first, then explicit registered defects are
injected. This is fixture-generation code, not the production cleaning pipeline.
"""

from __future__ import annotations

import copy
import csv
import hashlib
import json
import random
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


SEED = 1007
DEMO_DATE = date(2026, 10, 7)
TIMEZONE = "Europe/Paris"
SPECIFICATION_VERSION = "1.0.1"

SALES_FIELDS = ["order_id", "customer_id", "product_id", "order_date", "quantity", "unit_price_eur", "country"]
CUSTOMER_FIELDS = ["customer_id", "customer_name", "email", "country"]
PRODUCT_FIELDS = ["product_id", "product_name", "category", "unit_price_eur"]

RAW_FILES = [
    "sales_2026_07.csv", "sales_2026_08.xlsx", "sales_2026_09.csv",
    "sales_2026_10.xlsx", "customers.xlsx", "products.csv",
]
EXPECTED_FILES = [
    "ground_truth_sales.csv", "ground_truth_customers.csv",
    "ground_truth_products.csv", "expected_transactions.csv",
    "expected_outcomes.json",
]

SALES_LAYOUTS = {
    "sales_2026_07.csv": {
        "sheet": "",
        "headers": [("Order ID", "order_id"), ("Customer ID", "customer_id"), ("Product ID", "product_id"), ("Order Date", "order_date"), ("Quantity", "quantity"), ("Unit Price", "unit_price_eur"), ("Country", "country"), ("sales_channel", "_unexpected")],
    },
    "sales_2026_08.xlsx": {
        "sheet": "sales",
        "headers": [("OrderID", "order_id"), ("CustomerID", "customer_id"), ("ProductID", "product_id"), ("Date", "order_date"), ("Qty", "quantity"), ("Price EUR", "unit_price_eur"), ("Country Code", "country"), ("legacy_note", "_unexpected")],
    },
    "sales_2026_09.csv": {
        "sheet": "",
        "headers": [("order id", "order_id"), ("customer_id", "customer_id"), ("product id", "product_id"), ("order date", "order_date"), ("quantity", "quantity"), ("unit price", "unit_price_eur"), ("country code", "country"), ("campaign", "_unexpected")],
    },
    "sales_2026_10.xlsx": {
        "sheet": "sales",
        "headers": [("order_id", "order_id"), ("customer id", "customer_id"), ("product_id", "product_id"), ("order_date", "order_date"), ("quantity", "quantity"), ("unit_price_eur", "unit_price_eur"), ("country", "country"), ("review_note", "_unexpected")],
    },
}

CASE_DEFINITIONS = {
    "A_VALID": ("A", "Clean valid transaction."),
    "B_FORMATTING": ("B", "Recoverable formatting only; normalize, audit, accept."),
    "C_INVALID_QUANTITY": ("C", "Zero quantity; quarantine."),
    "D_EXACT_DUPLICATE": ("D", "Exact valid duplicate; first survives, copy deduplicates."),
    "E_CONFLICTING_ORDER_ID": ("E", "Conflicting valid order key; quarantine every member."),
    "F_HARD_INVALID_BEFORE_DUPLICATE": ("F", "Hard-invalid row quarantines before same-key duplicate accounting."),
    "G_FUZZY_CUSTOMER": ("G", "Similar customer names are flagged and never merged."),
    "H_UNKNOWN_CUSTOMER": ("H", "Unknown customer reference; quarantine transaction."),
    "H_UNKNOWN_PRODUCT": ("H", "Unknown product reference; quarantine transaction."),
    "I_FUTURE_DATE": ("I", "Date after frozen clock; quarantine."),
    "J_MULTIPLE_HARD_FAILURES": ("J", "One quarantine disposition with all hard failures evidenced."),
    "DATE_IMPOSSIBLE": (None, "Impossible calendar date."),
    "CURRENCY_NEGATIVE": (None, "Negative transaction price."),
    "CURRENCY_AMBIGUOUS": (None, "Ambiguous currency separator."),
    "MISSING_PRODUCT_ID": (None, "Missing critical product identifier."),
    "OPTIONAL_COUNTRY_MISSING": (None, "Optional country is null and retained."),
    "COUNTRY_FRANCE": (None, "Country alias France."),
    "COUNTRY_FRANCE_LOWER": (None, "Country alias france."),
    "COUNTRY_FR": (None, "Canonical country FR."),
    "COUNTRY_FRA": (None, "Country alias FRA."),
    "COUNTRY_UNKNOWN": (None, "Unknown country retained and flagged."),
    "DATE_ISO": (None, "ISO date form."),
    "DATE_DDMMYYYY": (None, "Day-first numeric date form."),
    "DATE_TEXTUAL": (None, "Bounded English textual date form."),
    "CURRENCY_EU": (None, "European grouped EUR form."),
    "CURRENCY_US": (None, "US grouped EUR form."),
    "CURRENCY_PLAIN": (None, "Plain integer EUR form."),
    "TEXT_HYGIENE": (None, "Identifier whitespace and case normalization."),
    "MISSING_ORDER_ID": (None, "Missing critical order identifier."),
    "NORMALIZED_DUPLICATE": (None, "Different raw forms normalize to one legal duplicate."),
    "INVALID_IDENTIFIER": (None, "Identifier outside the bounded pattern."),
    "FRACTIONAL_QUANTITY": (None, "Fractional quantity violates integer rule."),
    "MISSING_UNIT_PRICE": (None, "Missing critical transaction price."),
    "CUSTOMER_TEXT_HYGIENE": (None, "Customer name whitespace normalization."),
    "EMAIL_HYGIENE": (None, "Recoverable email trim/domain-case normalization."),
    "EMAIL_MALFORMED": (None, "Malformed customer email retained and flagged."),
    "CUSTOMER_OPTIONAL_NAME_MISSING": (None, "Optional customer name is null."),
    "CUSTOMER_COUNTRY_UNKNOWN": (None, "Unknown customer country retained and flagged."),
    "CUSTOMER_REFERENCE_MISSING_ID": (None, "Customer reference missing critical key."),
    "CUSTOMER_REFERENCE_DUPLICATE": (None, "Exact customer reference duplicate."),
    "CATEGORY_TECH": (None, "Tech category alias."),
    "CATEGORY_HOME_LIVING": (None, "Home & Living category alias."),
    "CATEGORY_OFFICE_SUPPLIES": (None, "Office Supplies category alias."),
    "CATEGORY_CLOTHING": (None, "Clothing category alias."),
    "CATEGORY_UNKNOWN": (None, "Unknown category retained and flagged."),
    "PRODUCT_OPTIONAL_NAME_MISSING": (None, "Optional product name is null."),
    "PRODUCT_REFERENCE_MISSING_ID": (None, "Product reference missing critical key."),
    "PRODUCT_REFERENCE_INVALID_PRICE": (None, "Product reference has negative price."),
    "PRODUCT_REFERENCE_DUPLICATE": (None, "Exact product reference duplicate."),
    "SCHEMA_ALIASES": (None, "Documented bounded header aliases across monthly files."),
    "SCHEMA_UNEXPECTED_COLUMNS": (None, "Unexpected columns are present for reporting."),
    "SCHEMA_UNEXPECTED_SHEET": (None, "One workbook contains an unselected notes sheet."),
}


def _customer_ground_truth() -> list[dict[str, Any]]:
    names = [
        "Aline Martin", "Bruno Leroy", "Camille Roux", "Nora Bernard", "Hugo Petit", "Lea Moreau",
        "Louis Fournier", "Emma Girard", "Theo Andre", "Ines Mercier", "Jules Dupont", "Chloe Lambert",
        "Noah Bonnet", "Lina Francois", "Adam Fontaine", "Eva Robin", "Leo Clement", "Maya Gauthier",
        "Sami Perrin", "Mila Durant", "Mila Durand", "Anna Renard", "Paul Marchand", "Sara Blanchard",
    ]
    return [
        {"customer_id": f"C{index:03d}", "customer_name": name, "email": f"{name.lower().replace(' ', '.')}@example.test", "country": "FR"}
        for index, name in enumerate(names, 1)
    ]


def _product_ground_truth() -> list[dict[str, Any]]:
    values = [
        ("Laptop Stand", "ELECTRONICS", "49.90"), ("Wireless Mouse", "ELECTRONICS", "29.95"),
        ("Desk Lamp", "HOME", "39.50"), ("Notebook Set", "OFFICE", "12.00"),
        ("Cotton T-Shirt", "APPAREL", "24.99"), ("Monitor", "ELECTRONICS", "299.00"),
        ("Storage Box", "HOME", "18.50"), ("Fountain Pen", "OFFICE", "15.75"),
        ("Hoodie", "APPAREL", "59.00"), ("Headphones", "ELECTRONICS", "79.99"),
        ("Coffee Maker", "HOME", "129.90"), ("Office Chair", "OFFICE", "1299.00"),
    ]
    return [
        {"product_id": f"P{index:03d}", "product_name": name, "category": category, "unit_price_eur": price}
        for index, (name, category, price) in enumerate(values, 1)
    ]


def _sales_ground_truth(customers: list[dict[str, Any]], products: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rng = random.Random(SEED)
    prices = {row["product_id"]: row["unit_price_eur"] for row in products}
    rows: list[dict[str, Any]] = []
    ordinal = 0
    for month, day_limit in ((7, 28), (8, 28), (9, 28), (10, 7)):
        for _ in range(40):
            ordinal += 1
            customer = customers[rng.randrange(len(customers))]["customer_id"]
            product = products[rng.randrange(len(products))]["product_id"]
            rows.append({
                "order_id": f"ORD-{ordinal:04d}", "customer_id": customer, "product_id": product,
                "order_date": date(2026, month, rng.randint(1, day_limit)).isoformat(),
                "quantity": rng.randint(1, 5), "unit_price_eur": prices[product], "country": "FR",
            })
    overrides = {
        "ORD-0001": ("C001", "P001", "2026-07-01", 1), "ORD-0002": ("C002", "P012", "2026-07-02", 2),
        "ORD-0003": ("C003", "P003", "2026-07-03", 2), "ORD-0004": ("C004", "P004", "2026-07-04", 1),
        "ORD-0005": ("C005", "P005", "2026-07-05", 2), "ORD-0006": ("C006", "P006", "2026-07-06", 2),
        "ORD-0007": ("C020", "P007", "2026-07-07", 1), "ORD-0030": ("C010", "P012", "2026-07-28", 1),
        "ORD-0121": ("C011", "P011", "2026-10-06", 1), "ORD-0122": ("C012", "P012", "2026-10-07", 2),
    }
    by_id = {row["order_id"]: row for row in rows}
    for order_id, (customer, product, order_date, quantity) in overrides.items():
        by_id[order_id].update({"customer_id": customer, "product_id": product, "order_date": order_date, "quantity": quantity, "unit_price_eur": prices[product]})
    for order_id in ("ORD-0025", "ORD-0026", "ORD-0027"):
        by_id[order_id].update({"product_id": "P012", "unit_price_eur": "1299.00"})
    return rows


def _fixture(row: dict[str, Any], fixture_id: str, *, case_ids: Iterable[str] = ()) -> dict[str, Any]:
    result = copy.deepcopy(row)
    result.update({
        "_fixture_id": fixture_id, "_ground_truth_order_id": row.get("order_id", ""),
        "_case_ids": list(case_ids), "_defects": [], "_expected_codes": [],
        "_expected_disposition": "ACCEPTED", "_expected_normalization": {},
        "_related_fixture_id": "", "_unexpected": "web",
    })
    return result


def _inject_sales_defects(ground_truth: list[dict[str, Any]]) -> dict[str, list[dict[str, Any]]]:
    records = {row["order_id"]: _fixture(row, f"TX-{index:04d}") for index, row in enumerate(ground_truth, 1)}

    def configure(order_id: str, *, cases: Iterable[str] = (), changes: dict[str, Any] | None = None,
                  defects: Iterable[str] = (), codes: Iterable[str] = (), disposition: str = "ACCEPTED",
                  normalization: dict[str, Any] | None = None) -> dict[str, Any]:
        record = records[order_id]
        record["_case_ids"].extend(cases)
        record["_defects"].extend(defects)
        record["_expected_codes"].extend(codes)
        record["_expected_disposition"] = disposition
        record["_expected_normalization"].update(normalization or {})
        record.update(changes or {})
        return record

    configure("ORD-0001", cases=["A_VALID"])
    configure("ORD-0002", cases=["B_FORMATTING"], changes={"order_id": " ord-0002 ", "customer_id": " c002 ", "product_id": "p012", "order_date": "02/07/2026", "quantity": "2", "unit_price_eur": "1 299,00 EUR", "country": " france "}, defects=["identifier_case_whitespace", "date_format", "currency_locale", "country_alias"], codes=["NORMALIZE_IDENTIFIER", "NORMALIZE_DATE", "NORMALIZE_CURRENCY", "NORMALIZE_COUNTRY"], normalization={"order_id": "ORD-0002", "customer_id": "C002", "product_id": "P012", "order_date": "2026-07-02", "quantity": 2, "unit_price_eur": "1299.00", "country": "FR"})
    configure("ORD-0003", cases=["C_INVALID_QUANTITY"], changes={"quantity": 0}, defects=["quantity_nonpositive"], codes=["INVALID_QUANTITY"], disposition="QUARANTINED")
    configure("ORD-0004", cases=["D_EXACT_DUPLICATE"])
    configure("ORD-0005", cases=["E_CONFLICTING_ORDER_ID"], defects=["duplicate_key_conflict"], codes=["DUPLICATE_KEY_CONFLICT"], disposition="QUARANTINED")
    configure("ORD-0006", cases=["F_HARD_INVALID_BEFORE_DUPLICATE"], changes={"quantity": -2}, defects=["quantity_nonpositive", "duplicate_like_after_hard_failure"], codes=["INVALID_QUANTITY"], disposition="QUARANTINED")
    configure("ORD-0007", cases=["G_FUZZY_CUSTOMER"])
    configure("ORD-0008", cases=["H_UNKNOWN_CUSTOMER"], changes={"customer_id": "C999"}, defects=["unknown_customer_reference"], codes=["UNKNOWN_CUSTOMER_REFERENCE"], disposition="QUARANTINED")
    configure("ORD-0009", cases=["H_UNKNOWN_PRODUCT"], changes={"product_id": "P999"}, defects=["unknown_product_reference"], codes=["UNKNOWN_PRODUCT_REFERENCE"], disposition="QUARANTINED")
    configure("ORD-0012", cases=["DATE_IMPOSSIBLE"], changes={"order_date": "31/02/2026"}, defects=["impossible_date"], codes=["INVALID_ORDER_DATE"], disposition="QUARANTINED")
    configure("ORD-0013", cases=["CURRENCY_NEGATIVE"], changes={"unit_price_eur": "-5.00"}, defects=["negative_unit_price"], codes=["INVALID_UNIT_PRICE"], disposition="QUARANTINED")
    configure("ORD-0014", cases=["CURRENCY_AMBIGUOUS"], changes={"unit_price_eur": "1,299"}, defects=["ambiguous_currency"], codes=["INVALID_UNIT_PRICE"], disposition="QUARANTINED")
    configure("ORD-0015", cases=["MISSING_PRODUCT_ID"], changes={"product_id": ""}, defects=["missing_critical_identifier"], codes=["MISSING_PRODUCT_ID"], disposition="QUARANTINED")
    configure("ORD-0016", cases=["OPTIONAL_COUNTRY_MISSING"], changes={"country": None}, defects=["missing_optional_field"], normalization={"country": None})
    configure("ORD-0017", cases=["COUNTRY_FRANCE"], changes={"country": "France"}, defects=["country_alias"], codes=["NORMALIZE_COUNTRY"], normalization={"country": "FR"})
    configure("ORD-0018", cases=["COUNTRY_FRANCE_LOWER"], changes={"country": "france"}, defects=["country_alias"], codes=["NORMALIZE_COUNTRY"], normalization={"country": "FR"})
    configure("ORD-0019", cases=["COUNTRY_FR"], changes={"country": "FR"}, normalization={"country": "FR"})
    configure("ORD-0020", cases=["COUNTRY_FRA"], changes={"country": "FRA"}, defects=["country_alias"], codes=["NORMALIZE_COUNTRY"], normalization={"country": "FR"})
    configure("ORD-0021", cases=["COUNTRY_UNKNOWN"], changes={"country": "Atlantis"}, defects=["unknown_country"], codes=["UNKNOWN_COUNTRY"], normalization={"country": "Atlantis"})
    configure("ORD-0022", cases=["DATE_ISO"], changes={"order_date": "2026-07-22"}, normalization={"order_date": "2026-07-22"})
    configure("ORD-0023", cases=["DATE_DDMMYYYY"], changes={"order_date": "23/07/2026"}, defects=["date_format"], codes=["NORMALIZE_DATE"], normalization={"order_date": "2026-07-23"})
    configure("ORD-0024", cases=["DATE_TEXTUAL"], changes={"order_date": "24 July 2026"}, defects=["date_format"], codes=["NORMALIZE_DATE"], normalization={"order_date": "2026-07-24"})
    for order_id, case_id, raw_price in (("ORD-0025", "CURRENCY_EU", "1 299,00 EUR"), ("ORD-0026", "CURRENCY_US", "EUR1,299.00"), ("ORD-0027", "CURRENCY_PLAIN", "1299")):
        configure(order_id, cases=[case_id], changes={"unit_price_eur": raw_price}, defects=["currency_locale"], codes=["NORMALIZE_CURRENCY"], normalization={"unit_price_eur": "1299.00"})
    configure("ORD-0028", cases=["TEXT_HYGIENE"], changes={"order_id": " ord-0028 ", "customer_id": " c008 ", "product_id": " p008 "}, defects=["identifier_case_whitespace"], codes=["NORMALIZE_IDENTIFIER"], normalization={"order_id": "ORD-0028", "customer_id": "C008", "product_id": "P008"})
    configure("ORD-0029", cases=["MISSING_ORDER_ID"], changes={"order_id": ""}, defects=["missing_critical_identifier"], codes=["MISSING_ORDER_ID"], disposition="QUARANTINED")
    configure("ORD-0030", cases=["NORMALIZED_DUPLICATE"])
    configure("ORD-0031", cases=["INVALID_IDENTIFIER"], changes={"customer_id": "BAD ID!"}, defects=["invalid_identifier"], codes=["INVALID_IDENTIFIER"], disposition="QUARANTINED")
    configure("ORD-0032", cases=["FRACTIONAL_QUANTITY"], changes={"quantity": "2.5"}, defects=["fractional_quantity"], codes=["INVALID_QUANTITY"], disposition="QUARANTINED")
    configure("ORD-0035", cases=["MISSING_UNIT_PRICE"], changes={"unit_price_eur": ""}, defects=["missing_unit_price"], codes=["INVALID_UNIT_PRICE"], disposition="QUARANTINED")
    configure("ORD-0121", cases=["I_FUTURE_DATE"], changes={"order_date": "2026-10-08"}, defects=["future_date"], codes=["FUTURE_ORDER_DATE"], disposition="QUARANTINED")
    configure("ORD-0122", cases=["J_MULTIPLE_HARD_FAILURES"], changes={"customer_id": "", "order_date": "09/10/2026", "quantity": -3}, defects=["missing_critical_identifier", "future_date", "quantity_nonpositive"], codes=["MISSING_CUSTOMER_ID", "FUTURE_ORDER_DATE", "INVALID_QUANTITY"], disposition="QUARANTINED")

    by_month: dict[str, list[dict[str, Any]]] = {name: [] for name in SALES_LAYOUTS}
    for index, truth in enumerate(ground_truth, 1):
        month = 7 + (index - 1) // 40
        filename = f"sales_2026_{month:02d}.csv" if month in (7, 9) else f"sales_2026_{month:02d}.xlsx"
        records[truth["order_id"]]["_unexpected"] = "web" if month in (7, 9) else "legacy"
        by_month[filename].append(records[truth["order_id"]])

    july = by_month["sales_2026_07.csv"]
    exact = copy.deepcopy(records["ORD-0004"])
    exact.update({"_fixture_id": "TX-D-DUP", "_expected_disposition": "DEDUPLICATED", "_expected_codes": ["DUPLICATE_EXACT"], "_defects": ["exact_row_duplicate"], "_related_fixture_id": records["ORD-0004"]["_fixture_id"]})
    july.append(exact)
    conflict = copy.deepcopy(records["ORD-0005"])
    conflict.update({"_fixture_id": "TX-E-CONFLICT", "quantity": int(conflict["quantity"]) + 1, "_related_fixture_id": records["ORD-0005"]["_fixture_id"]})
    july.append(conflict)
    valid = _fixture(next(row for row in ground_truth if row["order_id"] == "ORD-0006"), "TX-F-VALID", case_ids=["F_HARD_INVALID_BEFORE_DUPLICATE"])
    valid.update({"_unexpected": records["ORD-0006"]["_unexpected"], "_related_fixture_id": records["ORD-0006"]["_fixture_id"]})
    july.append(valid)
    normalized = copy.deepcopy(records["ORD-0030"])
    normalized.update({"_fixture_id": "TX-NORMALIZED-DUP", "order_id": " ord-0030 ", "customer_id": " c010 ", "product_id": "p012", "order_date": "28/07/2026", "quantity": "1", "unit_price_eur": "1 299,00 EUR", "country": "France", "_defects": ["semantic_duplicate_after_normalization"], "_expected_codes": ["NORMALIZE_IDENTIFIER", "NORMALIZE_DATE", "NORMALIZE_CURRENCY", "NORMALIZE_COUNTRY", "DUPLICATE_EXACT"], "_expected_disposition": "DEDUPLICATED", "_expected_normalization": {"order_id": "ORD-0030", "customer_id": "C010", "product_id": "P012", "order_date": "2026-07-28", "quantity": 1, "unit_price_eur": "1299.00", "country": "FR"}, "_related_fixture_id": records["ORD-0030"]["_fixture_id"]})
    july.append(normalized)
    return by_month


def _inject_customer_defects(ground_truth: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows = [_fixture(row, f"CUSTOMER-{index:03d}") for index, row in enumerate(ground_truth, 1)]
    for row in rows:
        row["_expected_disposition"] = "REFERENCE_VALID"
    by_id = {row["customer_id"]: row for row in rows}
    by_id["C002"].update({"customer_name": "  Bruno   Leroy  ", "country": "france", "_case_ids": ["CUSTOMER_TEXT_HYGIENE"], "_defects": ["text_whitespace", "country_alias"], "_expected_codes": ["NORMALIZE_WHITESPACE", "NORMALIZE_COUNTRY"], "_expected_normalization": {"customer_name": "Bruno Leroy", "country": "FR"}})
    by_id["C003"].update({"email": " camille.roux@EXAMPLE.TEST ", "country": "FRA", "_case_ids": ["EMAIL_HYGIENE"], "_defects": ["email_hygiene", "country_alias"], "_expected_codes": ["NORMALIZE_EMAIL", "NORMALIZE_COUNTRY"], "_expected_normalization": {"email": "camille.roux@example.test", "country": "FR"}})
    by_id["C004"].update({"email": "nora.bernard.example.test", "_case_ids": ["EMAIL_MALFORMED"], "_defects": ["malformed_email"], "_expected_codes": ["INVALID_EMAIL_FORMAT"]})
    by_id["C005"].update({"customer_name": None, "_case_ids": ["CUSTOMER_OPTIONAL_NAME_MISSING"], "_defects": ["missing_optional_field"]})
    by_id["C006"].update({"country": "Atlantis", "_case_ids": ["CUSTOMER_COUNTRY_UNKNOWN"], "_defects": ["unknown_country"], "_expected_codes": ["UNKNOWN_COUNTRY"]})
    for customer_id in ("C020", "C021"):
        by_id[customer_id]["_case_ids"] = ["G_FUZZY_CUSTOMER"]
        by_id[customer_id]["_expected_codes"] = ["FUZZY_CUSTOMER_CANDIDATE"]
    missing = _fixture({"customer_id": "", "customer_name": "Synthetic Missing ID", "email": "missing.id@example.test", "country": "FR"}, "CUSTOMER-MISSING-ID", case_ids=["CUSTOMER_REFERENCE_MISSING_ID"])
    missing.update({"_defects": ["missing_critical_identifier"], "_expected_codes": ["MISSING_CUSTOMER_ID"], "_expected_disposition": "REFERENCE_REJECTED"})
    rows.append(missing)
    duplicate = copy.deepcopy(by_id["C010"])
    duplicate.update({"_fixture_id": "CUSTOMER-EXACT-DUP", "_case_ids": ["CUSTOMER_REFERENCE_DUPLICATE"], "_defects": ["exact_reference_duplicate"], "_expected_codes": ["REFERENCE_DUPLICATE_EXACT"], "_expected_disposition": "REFERENCE_DEDUPLICATED", "_related_fixture_id": by_id["C010"]["_fixture_id"]})
    rows.append(duplicate)
    return rows


def _inject_product_defects(ground_truth: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows = [_fixture(row, f"PRODUCT-{index:03d}") for index, row in enumerate(ground_truth, 1)]
    for row in rows:
        row["_expected_disposition"] = "REFERENCE_VALID"
    by_id = {row["product_id"]: row for row in rows}
    for product_id, raw, case_id, canonical in (("P001", "Tech", "CATEGORY_TECH", "ELECTRONICS"), ("P003", "Home & Living", "CATEGORY_HOME_LIVING", "HOME"), ("P004", "office supplies", "CATEGORY_OFFICE_SUPPLIES", "OFFICE"), ("P005", "Clothing", "CATEGORY_CLOTHING", "APPAREL")):
        by_id[product_id].update({"category": raw, "_case_ids": [case_id], "_defects": ["category_alias"], "_expected_codes": ["NORMALIZE_CATEGORY"], "_expected_normalization": {"category": canonical}})
    by_id["P006"].update({"category": "Gadgets", "_case_ids": ["CATEGORY_UNKNOWN"], "_defects": ["unknown_category"], "_expected_codes": ["UNKNOWN_CATEGORY"]})
    by_id["P007"].update({"product_name": None, "_case_ids": ["PRODUCT_OPTIONAL_NAME_MISSING"], "_defects": ["missing_optional_field"]})
    by_id["P012"].update({"unit_price_eur": "1 299,00 EUR", "_case_ids": ["CURRENCY_EU"], "_defects": ["currency_locale"], "_expected_codes": ["NORMALIZE_CURRENCY"], "_expected_normalization": {"unit_price_eur": "1299.00"}})
    missing = _fixture({"product_id": "", "product_name": "Synthetic Missing SKU", "category": "OFFICE", "unit_price_eur": "10.00"}, "PRODUCT-MISSING-ID", case_ids=["PRODUCT_REFERENCE_MISSING_ID"])
    missing.update({"_defects": ["missing_critical_identifier"], "_expected_codes": ["MISSING_PRODUCT_ID"], "_expected_disposition": "REFERENCE_REJECTED"})
    rows.append(missing)
    invalid = _fixture({"product_id": "P013", "product_name": "Synthetic Invalid Price", "category": "OFFICE", "unit_price_eur": "-3.00"}, "PRODUCT-INVALID-PRICE", case_ids=["PRODUCT_REFERENCE_INVALID_PRICE"])
    invalid.update({"_defects": ["negative_unit_price"], "_expected_codes": ["INVALID_UNIT_PRICE"], "_expected_disposition": "REFERENCE_REJECTED"})
    rows.append(invalid)
    duplicate = copy.deepcopy(by_id["P010"])
    duplicate.update({"_fixture_id": "PRODUCT-EXACT-DUP", "_case_ids": ["PRODUCT_REFERENCE_DUPLICATE"], "_defects": ["exact_reference_duplicate"], "_expected_codes": ["REFERENCE_DUPLICATE_EXACT"], "_expected_disposition": "REFERENCE_DEDUPLICATED", "_related_fixture_id": by_id["P010"]["_fixture_id"]})
    rows.append(duplicate)
    return rows


def _write_csv(path: Path, headers: list[tuple[str, str]], rows: list[dict[str, Any]]) -> None:
    with path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.writer(stream, lineterminator="\n")
        writer.writerow([label for label, _ in headers])
        for row in rows:
            writer.writerow(["" if row.get(key) is None else row.get(key, "") for _, key in headers])


def _write_xlsx(path: Path, sheet_name: str, headers: list[tuple[str, str]], rows: list[dict[str, Any]], *, add_unexpected_sheet: bool = False) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    labels = [label for label, _ in headers]
    sheet.append(labels)
    for row in rows:
        sheet.append([row.get(key) for _, key in headers])
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.fill = fill
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="center")
    for index, (label, key) in enumerate(headers, 1):
        values = [label] + ["" if row.get(key) is None else str(row.get(key, "")) for row in rows]
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = min(max(map(len, values)) + 2, 28)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    if add_unexpected_sheet:
        notes = workbook.create_sheet("notes")
        notes["A1"] = "Synthetic demo note: this sheet is intentionally not selected."
        notes["A1"].font = Font(name="Arial", size=10)
        notes.column_dimensions["A"].width = 62
        notes.sheet_view.showGridLines = False
    fixed = datetime(2026, 10, 7)
    workbook.properties.creator = "DataForge deterministic demo generator"
    workbook.properties.created = fixed
    workbook.properties.modified = fixed
    workbook.save(path)


def _assign_source(rows: list[dict[str, Any]], filename: str, sheet: str) -> None:
    for source_row, row in enumerate(rows, 2):
        row.update({"_source_file": filename, "_source_sheet": sheet, "_source_row": source_row})


def _source(record: dict[str, Any], record_type: str) -> dict[str, Any]:
    return {"fixture_id": record["_fixture_id"], "record_type": record_type, "source_file": record["_source_file"], "source_sheet": record["_source_sheet"], "source_row": record["_source_row"]}


def _semantic_xlsx_sha256(path: Path) -> str:
    workbook = load_workbook(path, data_only=False, read_only=True)
    payload = [{"sheet": sheet.title, "rows": [[cell.value for cell in row] for row in sheet.iter_rows()]} for sheet in workbook.worksheets]
    workbook.close()
    encoded = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":"), default=str).encode()
    return hashlib.sha256(encoded).hexdigest()


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _case_manifest(all_records: list[tuple[str, dict[str, Any]]]) -> list[dict[str, Any]]:
    cases = []
    for case_id, (interaction_case, description) in CASE_DEFINITIONS.items():
        matching = [(record_type, record) for record_type, record in all_records if case_id in record.get("_case_ids", [])]
        sources = [_source(record, record_type) for record_type, record in matching]
        if case_id == "SCHEMA_ALIASES":
            sources = [{"record_type": "sales", "source_file": name, "source_sheet": layout["sheet"], "headers": [label for label, _ in layout["headers"]]} for name, layout in SALES_LAYOUTS.items()] + [{"record_type": "customers", "source_file": "customers.xlsx", "source_sheet": "customers", "headers": ["Customer ID", "Customer Name", "Email Address", "Country Code", "legacy_segment"]}, {"record_type": "products", "source_file": "products.csv", "source_sheet": "", "headers": ["SKU", "Product Name", "Product Category", "Price EUR", "supplier_note"]}]
        elif case_id == "SCHEMA_UNEXPECTED_COLUMNS":
            sources = [{"source_file": name, "unexpected_header": layout["headers"][-1][0]} for name, layout in SALES_LAYOUTS.items()]
        elif case_id == "SCHEMA_UNEXPECTED_SHEET":
            sources = [{"source_file": "sales_2026_08.xlsx", "source_sheet": "notes"}]
        codes: list[str] = []
        defects: list[str] = []
        normalizations = []
        dispositions = []
        relationships = []
        for _, record in matching:
            for value in record["_expected_codes"]:
                if value not in codes:
                    codes.append(value)
            for value in record["_defects"]:
                if value not in defects:
                    defects.append(value)
            if record["_expected_normalization"]:
                normalizations.append({"fixture_id": record["_fixture_id"], "values": record["_expected_normalization"]})
            dispositions.append(record["_expected_disposition"])
            if record["_related_fixture_id"]:
                relationships.append({"fixture_id": record["_fixture_id"], "related_fixture_id": record["_related_fixture_id"]})
        cases.append({"case_id": case_id, "interaction_case": interaction_case, "description": description, "source_records": sources, "injected_defects": defects, "expected_normalization": normalizations, "expected_codes": codes, "expected_terminal_dispositions": dispositions, "relationships": relationships})
    return cases


def _coverage() -> dict[str, list[str]]:
    return {
        "schema.documented_aliases": ["SCHEMA_ALIASES"], "schema.unexpected_columns": ["SCHEMA_UNEXPECTED_COLUMNS"], "schema.unexpected_sheet": ["SCHEMA_UNEXPECTED_SHEET"],
        "locale.currency_formats": ["B_FORMATTING", "CURRENCY_EU", "CURRENCY_US", "CURRENCY_PLAIN"], "locale.currency_invalid": ["CURRENCY_NEGATIVE", "CURRENCY_AMBIGUOUS", "MISSING_UNIT_PRICE"],
        "dates.iso": ["DATE_ISO"], "dates.ddmmyyyy": ["B_FORMATTING", "DATE_DDMMYYYY"], "dates.textual": ["DATE_TEXTUAL"], "dates.impossible": ["DATE_IMPOSSIBLE"], "dates.future": ["I_FUTURE_DATE", "J_MULTIPLE_HARD_FAILURES"],
        "country.aliases": ["COUNTRY_FRANCE", "COUNTRY_FRANCE_LOWER", "COUNTRY_FR", "COUNTRY_FRA"], "country.unknown": ["COUNTRY_UNKNOWN", "CUSTOMER_COUNTRY_UNKNOWN"],
        "categories.aliases": ["CATEGORY_TECH", "CATEGORY_HOME_LIVING", "CATEGORY_OFFICE_SUPPLIES", "CATEGORY_CLOTHING"], "categories.unknown": ["CATEGORY_UNKNOWN"],
        "missingness.critical_identifier": ["MISSING_ORDER_ID", "MISSING_PRODUCT_ID", "J_MULTIPLE_HARD_FAILURES", "CUSTOMER_REFERENCE_MISSING_ID", "PRODUCT_REFERENCE_MISSING_ID"], "missingness.optional_field": ["OPTIONAL_COUNTRY_MISSING", "CUSTOMER_OPTIONAL_NAME_MISSING", "PRODUCT_OPTIONAL_NAME_MISSING"],
        "duplicates.exact_row": ["D_EXACT_DUPLICATE"], "duplicates.normalized_business_key": ["NORMALIZED_DUPLICATE"], "duplicates.conflicting_business_key": ["E_CONFLICTING_ORDER_ID"], "duplicates.fuzzy_customer": ["G_FUZZY_CUSTOMER"],
        "business.quantity_nonpositive": ["C_INVALID_QUANTITY", "F_HARD_INVALID_BEFORE_DUPLICATE", "J_MULTIPLE_HARD_FAILURES"], "business.quantity_fractional": ["FRACTIONAL_QUANTITY"], "business.negative_price": ["CURRENCY_NEGATIVE", "PRODUCT_REFERENCE_INVALID_PRICE"], "business.invalid_identifier": ["INVALID_IDENTIFIER"],
        "text.whitespace": ["B_FORMATTING", "TEXT_HYGIENE", "CUSTOMER_TEXT_HYGIENE"], "text.case": ["B_FORMATTING", "TEXT_HYGIENE", "EMAIL_HYGIENE"], "email.hygiene": ["EMAIL_HYGIENE"], "email.malformed": ["EMAIL_MALFORMED"],
        "references.valid": ["A_VALID"], "references.unknown_customer": ["H_UNKNOWN_CUSTOMER"], "references.unknown_product": ["H_UNKNOWN_PRODUCT"], "references.exact_duplicate": ["CUSTOMER_REFERENCE_DUPLICATE", "PRODUCT_REFERENCE_DUPLICATE"],
        "precedence.hard_before_duplicate": ["F_HARD_INVALID_BEFORE_DUPLICATE"], "failures.multiple": ["J_MULTIPLE_HARD_FAILURES"],
    }


def _write_expected_transactions(path: Path, rows: list[dict[str, Any]]) -> None:
    fields = ["fixture_id", "source_file", "source_sheet", "source_row", "ground_truth_order_id", "raw_order_id", "case_ids", "defects", "expected_codes", "expected_terminal_disposition", "related_fixture_id"]
    with path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields, lineterminator="\n")
        writer.writeheader()
        for row in rows:
            writer.writerow({"fixture_id": row["_fixture_id"], "source_file": row["_source_file"], "source_sheet": row["_source_sheet"], "source_row": row["_source_row"], "ground_truth_order_id": row["_ground_truth_order_id"], "raw_order_id": row.get("order_id", ""), "case_ids": "|".join(row["_case_ids"]), "defects": "|".join(row["_defects"]), "expected_codes": "|".join(row["_expected_codes"]), "expected_terminal_disposition": row["_expected_disposition"], "related_fixture_id": row["_related_fixture_id"]})


def _prepare_directories(root: Path) -> tuple[Path, Path]:
    root = root.resolve()
    raw_dir, expected_dir = root / "data" / "demo_raw", root / "data" / "demo_expected"
    if raw_dir.name != "demo_raw" or expected_dir.name != "demo_expected":
        raise ValueError("Refusing to write outside fixed demo directory names")
    raw_dir.mkdir(parents=True, exist_ok=True)
    expected_dir.mkdir(parents=True, exist_ok=True)
    for directory, filenames in ((raw_dir, RAW_FILES), (expected_dir, EXPECTED_FILES)):
        for filename in filenames:
            target = directory / filename
            if target.exists():
                target.unlink()
    return raw_dir, expected_dir


def generate_demo_data(root: Path) -> dict[str, Any]:
    """Generate the corpus below ``root/data`` and return its manifest."""
    raw_dir, expected_dir = _prepare_directories(root)
    customers_truth, products_truth = _customer_ground_truth(), _product_ground_truth()
    sales_truth = _sales_ground_truth(customers_truth, products_truth)
    sales_by_file = _inject_sales_defects(sales_truth)
    customer_rows, product_rows = _inject_customer_defects(customers_truth), _inject_product_defects(products_truth)

    for filename, rows in sales_by_file.items():
        layout = SALES_LAYOUTS[filename]
        _assign_source(rows, filename, layout["sheet"])
        path = raw_dir / filename
        if path.suffix == ".csv":
            _write_csv(path, layout["headers"], rows)
        else:
            _write_xlsx(path, layout["sheet"], layout["headers"], rows, add_unexpected_sheet=filename == "sales_2026_08.xlsx")

    customer_headers = [("Customer ID", "customer_id"), ("Customer Name", "customer_name"), ("Email Address", "email"), ("Country Code", "country"), ("legacy_segment", "_unexpected")]
    for row in customer_rows:
        row["_unexpected"] = "synthetic"
    _assign_source(customer_rows, "customers.xlsx", "customers")
    _write_xlsx(raw_dir / "customers.xlsx", "customers", customer_headers, customer_rows)
    product_headers = [("SKU", "product_id"), ("Product Name", "product_name"), ("Product Category", "category"), ("Price EUR", "unit_price_eur"), ("supplier_note", "_unexpected")]
    for row in product_rows:
        row["_unexpected"] = "synthetic"
    _assign_source(product_rows, "products.csv", "")
    _write_csv(raw_dir / "products.csv", product_headers, product_rows)

    _write_csv(expected_dir / "ground_truth_sales.csv", [(field, field) for field in SALES_FIELDS], sales_truth)
    _write_csv(expected_dir / "ground_truth_customers.csv", [(field, field) for field in CUSTOMER_FIELDS], customers_truth)
    _write_csv(expected_dir / "ground_truth_products.csv", [(field, field) for field in PRODUCT_FIELDS], products_truth)
    transactions = [row for filename in sorted(sales_by_file) for row in sales_by_file[filename]]
    _write_expected_transactions(expected_dir / "expected_transactions.csv", transactions)

    all_records = [("sales", row) for row in transactions] + [("customers", row) for row in customer_rows] + [("products", row) for row in product_rows]
    cases, coverage = _case_manifest(all_records), _coverage()
    counts = {disposition: sum(row["_expected_disposition"] == disposition for row in transactions) for disposition in ("ACCEPTED", "QUARANTINED", "DEDUPLICATED")}
    row_counts = {name: len(rows) for name, rows in sales_by_file.items()} | {"customers.xlsx": len(customer_rows), "products.csv": len(product_rows)}
    files = []
    for filename in RAW_FILES:
        path = raw_dir / filename
        integrity = {"comparison": "semantic", "semantic_sha256": _semantic_xlsx_sha256(path)} if path.suffix == ".xlsx" else {"comparison": "byte", "sha256": _sha256(path)}
        record_type = "customers" if filename == "customers.xlsx" else "products" if filename == "products.csv" else "sales"
        selected_sheets = ["customers"] if filename == "customers.xlsx" else ["sales"] if path.suffix == ".xlsx" else []
        files.append({"path": f"data/demo_raw/{filename}", "record_type": record_type, "format": path.suffix[1:], "selected_sheets": selected_sheets, "row_count": row_counts[filename], **integrity})
    for filename in EXPECTED_FILES[:-1]:
        path = expected_dir / filename
        files.append({"path": f"data/demo_expected/{filename}", "record_type": "expected_fixture", "format": path.suffix[1:], "comparison": "byte", "sha256": _sha256(path)})

    manifest = {
        "specification_version": SPECIFICATION_VERSION,
        "dataset": {"name": "DataForge deterministic synthetic e-commerce demo", "synthetic": True, "contains_real_customer_data": False, "seed": SEED, "demo_date": DEMO_DATE.isoformat(), "timezone": TIMEZONE, "determinism": {"csv_json": "byte-equivalent", "xlsx": "semantic record-equivalent; container bytes are not asserted"}},
        "counts": {"transaction_ground_truth": len(sales_truth), "transaction_input": len(transactions), "transaction_expected": {"accepted": counts["ACCEPTED"], "quarantined": counts["QUARANTINED"], "deduplicated": counts["DEDUPLICATED"]}, "customers_ground_truth": len(customers_truth), "customers_input": len(customer_rows), "products_ground_truth": len(products_truth), "products_input": len(product_rows)},
        "g2_expected": {"input_transaction_rows": len(transactions), "accepted_rows": counts["ACCEPTED"], "quarantined_rows": counts["QUARANTINED"], "deduplicated_rows": counts["DEDUPLICATED"], "reconciled": len(transactions) == sum(counts.values())},
        "files": files, "coverage": coverage, "cases": cases,
    }
    assert {case["interaction_case"] for case in cases if case["interaction_case"]} == set("ABCDEFGHIJ")
    assert all(coverage.values()) and manifest["g2_expected"]["reconciled"]
    assert (len(transactions), counts["ACCEPTED"], counts["QUARANTINED"], counts["DEDUPLICATED"]) == (164, 146, 16, 2)
    assert len({row["_fixture_id"] for row in transactions}) == len(transactions)
    assert all("@example.test" in row["email"].strip().lower() for row in customer_rows if row.get("email") and "@" in row["email"])
    (expected_dir / "expected_outcomes.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    return manifest


def main() -> None:
    manifest = generate_demo_data(Path(__file__).resolve().parents[1])
    counts = manifest["counts"]
    print(f"Generated deterministic synthetic demo data: {counts['transaction_input']} transactions, {counts['customers_input']} customers, {counts['products_input']} products (seed={SEED}).")


if __name__ == "__main__":
    main()
